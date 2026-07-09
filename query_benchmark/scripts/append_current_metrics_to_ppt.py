#!/usr/bin/env python3
"""Append current evaluation metrics/rules and visual figures to an existing PPT.

This script is intentionally presentation-oriented:
- adds concise protocol/rule slides (v0.4 style)
- appends selected visualization figures
- supports in-place overwrite with automatic backup
"""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def _pick_blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        name = str(getattr(layout, "name", "") or "").lower()
        if "blank" in name:
            return layout
    return prs.slide_layouts[len(prs.slide_layouts) - 1]


def _add_title(slide, title: str, subtitle: str | None = None) -> None:
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(1.0))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(22, 34, 56)
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Aptos"
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(93, 102, 119)


def _add_bullets(slide, x: float, y: float, w: float, h: float, bullets: list[str], title: str | None = None) -> None:
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Aptos"
        p0.font.bold = True
        p0.font.size = Pt(18)
        p0.font.color.rgb = RGBColor(31, 45, 70)
    else:
        tf.paragraphs[0].text = ""

    for i, line in enumerate(bullets):
        p = tf.add_paragraph() if (i > 0 or title) else tf.paragraphs[0]
        p.text = line
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(37, 50, 73)
        p.space_after = Pt(6)


def _add_table_slide(prs: Presentation, title: str, subtitle: str, headers: list[str], rows: list[list[str]]) -> None:
    slide = prs.slides.add_slide(_pick_blank_layout(prs))
    _add_title(slide, title, subtitle)

    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.45), Inches(1.35), Inches(12.4), Inches(5.7))
    table = table_shape.table

    col_widths = [2.2, 3.7, 6.5]
    if n_cols == 4:
        col_widths = [1.9, 2.9, 3.8, 3.8]
    for i, w in enumerate(col_widths[:n_cols]):
        table.columns[i].width = Inches(w)

    # Header
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Aptos"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(51, 95, 140)

    # Body
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Aptos"
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(34, 45, 67)
            p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            if r_idx % 2 == 1:
                cell.fill.fore_color.rgb = RGBColor(245, 248, 252)
            else:
                cell.fill.fore_color.rgb = RGBColor(233, 240, 249)


def _add_image_slide(prs: Presentation, title: str, image_path: Path, subtitle: str = "") -> None:
    slide = prs.slides.add_slide(_pick_blank_layout(prs))
    _add_title(slide, title, subtitle)
    slide.shapes.add_picture(str(image_path), Inches(0.35), Inches(1.15), width=Inches(12.65), height=Inches(5.85))


def run(args: argparse.Namespace) -> dict:
    ppt_in = args.ppt_in.expanduser().resolve()
    ppt_out = args.ppt_out.expanduser().resolve()
    if not ppt_in.exists():
        raise FileNotFoundError(f"PPT not found: {ppt_in}")

    backup_path = None
    if ppt_in == ppt_out:
        backup_path = ppt_in.with_name(f"{ppt_in.stem}.backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ppt_in.suffix}")
        shutil.copy2(ppt_in, backup_path)

    prs = Presentation(str(ppt_in))
    slide_count_before = len(prs.slides)

    # Slide 1: overview
    s = prs.slides.add_slide(_pick_blank_layout(prs))
    _add_title(
        s,
        "Current Evaluation Metrics (v0.4)",
        "Two blocks: Validation (deterministic integrity) + Analytics (family behavior)",
    )
    _add_bullets(
        s,
        x=0.7,
        y=1.55,
        w=11.8,
        h=5.5,
        title="Core principles",
        bullets=[
            "No forced single total score: keep channel-level transparency.",
            "Missing-as-value policy in subgroup/conditional families.",
            "Subgroup vs Tail split by support band: top 97% vs bottom 3%.",
            "Missingness family is N/A when real data has no native missing signal.",
            "Validation channels are deterministic and independently reportable.",
        ],
    )

    # Slide 2: validation table
    _add_table_slide(
        prs,
        title="Validation Metrics (Deterministic)",
        subtitle="Integrity-first checks; report channels separately",
        headers=["Channel", "Goal", "Rule Summary"],
        rows=[
            [
                "Cardinality / Range",
                "Value-space coverage",
                "Discrete: column-risk + value-missing ratio (50/50). Continuous: symmetric range penalty for under/over coverage.",
            ],
            [
                "Missing Introduction",
                "No new missing where real has none",
                "If violating columns >=2 -> hard fail. If exactly 1 -> base penalty + smooth penalty by missing rate.",
            ],
            [
                "Uniqueness Integrity",
                "Protect unique columns",
                "Categorical unique: duplicate introduces hard penalty. Numerical unique: duplicates <=10 tolerated; >10 hard fail.",
            ],
            [
                "Impossible State",
                "Constraint validity",
                "Independent channel; kept as a dedicated policy slot for invalid state checks.",
            ],
        ],
    )

    # Slide 3: analytics table
    _add_table_slide(
        prs,
        title="Analytics Metrics (Family-based)",
        subtitle="Family scores are independent; no forced analytics merge",
        headers=["Family", "Indicators", "Key Rules"],
        rows=[
            [
                "Subgroup Structure",
                "Internal profile + size stability (50/50)",
                "Build subgroup keys after missing tokenization; evaluate top 97% support groups.",
            ],
            [
                "Conditional Dependency",
                "Strength + direction + slice consistency",
                "Type-aware association estimators; include missing-involved slices when support gates pass.",
            ],
            [
                "Tail / Rarity",
                "Tail set + tail mass + concentration",
                "Tail = bottom 3% support keys; no default lower floor; head-protection cap can apply.",
            ],
            [
                "Missingness",
                "Marginal missing + co-missing (50/50)",
                "Enabled only when real data has native missing signal; otherwise return N/A.",
            ],
        ],
    )

    # Slide 4: operational checklist
    s4 = prs.slides.add_slide(_pick_blank_layout(prs))
    _add_title(s4, "Operational Checklist for Scoring Runs", "What should always be logged and reported")
    _add_bullets(
        s4,
        x=0.75,
        y=1.55,
        w=11.8,
        h=5.4,
        title="Run-time checklist",
        bullets=[
            "Record applicability gates (especially missingness N/A decisions).",
            "Preserve per-channel outputs instead of collapsing into one scalar.",
            "For subgroup/tail: keep support-band definition fixed and auditable.",
            "Log active indicators and renormalization when indicators are N/A.",
            "Export model-level tables and family/sub-item views for slide-ready diagnostics.",
        ],
    )

    # Visualization slides
    image_items = [
        ("C2 Capability Heatmap (Current Primary Run)", args.fig_c2_compact),
        ("M4 Capability Heatmap (Current Primary Run)", args.fig_m4_compact),
        ("N1 Capability Heatmap (Current Primary Run)", args.fig_n1_compact),
        ("C2 Family Sub-item Heatmap", args.fig_c2_subitems),
        ("M4 Family Sub-item Heatmap", args.fig_m4_subitems),
        ("N1 Family Sub-item Heatmap", args.fig_n1_subitems),
    ]
    for title, path in image_items:
        p = path.expanduser().resolve()
        if p.exists():
            _add_image_slide(prs, title, p, subtitle=str(p))

    prs.save(str(ppt_out))

    manifest = {
        "timestamp": datetime.now().isoformat(),
        "ppt_in": str(ppt_in),
        "ppt_out": str(ppt_out),
        "backup": str(backup_path) if backup_path else None,
        "slides_before": slide_count_before,
        "slides_after": len(prs.slides),
        "slides_added": len(prs.slides) - slide_count_before,
        "images_used": [str(path.expanduser().resolve()) for _, path in image_items if path.expanduser().resolve().exists()],
    }
    args.manifest_out.parent.mkdir(parents=True, exist_ok=True)
    args.manifest_out.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Append current evaluation metrics/rules + visuals to PPT.")
    parser.add_argument("--ppt-in", type=Path, required=True)
    parser.add_argument("--ppt-out", type=Path, required=True)
    parser.add_argument("--fig-c2-compact", type=Path, required=True)
    parser.add_argument("--fig-m4-compact", type=Path, required=True)
    parser.add_argument("--fig-n1-compact", type=Path, required=True)
    parser.add_argument("--fig-c2-subitems", type=Path, required=True)
    parser.add_argument("--fig-m4-subitems", type=Path, required=True)
    parser.add_argument("--fig-n1-subitems", type=Path, required=True)
    parser.add_argument(
        "--manifest-out",
        type=Path,
        default=Path("logs/analysis/ppt_v4_metrics/append_manifest.json"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
