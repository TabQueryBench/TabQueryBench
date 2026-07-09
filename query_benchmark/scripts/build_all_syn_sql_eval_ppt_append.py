#!/usr/bin/env python3
"""Build cross-dataset evaluation visuals and append them to an existing PPT.

Scope:
- Datasets that currently have both synthetic panel data and workload SQL runs.
- Uses latest experiment outputs (model_scores + selected_workloads) per dataset.
- Generates:
  1) per-dataset capability heatmaps (without overall column)
  2) per-metric cross-dataset heatmaps
  3) per-metric grouped bar charts
- Appends all generated figures to the end of a target PPT.
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]

MODEL_ORDER = [
    "realtabformer",
    "bayesnet",
    "arf",
    "tvae",
    "ctgan",
    "tabpfgen",
    "tabddpm",
]

MODEL_LABEL = {
    "realtabformer": "RealTabFormer",
    "bayesnet": "BayesNet",
    "arf": "ARF",
    "tvae": "TVAE",
    "ctgan": "CTGAN",
    "tabpfgen": "TabPFGen",
    "tabddpm": "TabDDPM",
}

METRICS_MAIN = [
    ("subgroup_structure_score", "Subgroup"),
    ("conditional_dependency_structure_score", "Conditional"),
    ("tail_rarity_structure_score", "Tail/Rarity"),
    ("missingness_structure_score", "Missingness"),
    ("validation_cardinality_range_score", "Cardinality"),
]

METRICS_ALL_SMALL = [
    ("subgroup_structure_score", "Subgroup"),
    ("conditional_dependency_structure_score", "Conditional"),
    ("tail_rarity_structure_score", "Tail/Rarity"),
    ("missingness_structure_score", "Missingness"),
    ("validation_cardinality_range_score", "Cardinality"),
    ("validation_missing_introduction_score", "Missing Intro"),
    ("validation_uniqueness_integrity_score", "Uniqueness"),
    ("query_success_rate", "Query Success"),
]


@dataclass
class DatasetScore:
    dataset_id: str
    experiment_dir: Path
    primary_workload_run_id: str
    # model_id -> metric_key -> value
    scores: dict[str, dict[str, float]]


def _custom_cmap() -> LinearSegmentedColormap:
    # Low -> high resembles the user-provided green style.
    colors = [
        "#a68a5b",
        "#9aae72",
        "#7fbf74",
        "#63b96d",
    ]
    return LinearSegmentedColormap.from_list("custom_green_brown", colors)


def _as_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None:
            return default
        return float(value)
    except Exception:
        return default


def _read_primary_scores(experiment_dir: Path, dataset_id: str) -> DatasetScore:
    selected_path = experiment_dir / f"selected_workloads_{dataset_id}.json"
    score_path = experiment_dir / f"model_scores_{dataset_id}.csv"

    selected = json.loads(selected_path.read_text(encoding="utf-8"))
    primary = str(selected.get("primary_workload_run_id") or "")

    rows: list[dict[str, str]] = []
    with score_path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if str(row.get("workload_run_id") or "") == primary:
                rows.append(row)

    scores: dict[str, dict[str, float]] = {}
    for row in rows:
        model_id = str(row.get("model_id") or "").strip().lower()
        if not model_id:
            continue
        m: dict[str, float] = {}
        for k in row.keys():
            if k in {"workload_run_id", "model_id", "n_repeats"}:
                continue
            m[k] = _as_float(row.get(k), 0.0)
        scores[model_id] = m

    return DatasetScore(
        dataset_id=dataset_id,
        experiment_dir=experiment_dir,
        primary_workload_run_id=primary,
        scores=scores,
    )


def _model_ids_present(ds_scores: list[DatasetScore]) -> list[str]:
    all_models: set[str] = set()
    for ds in ds_scores:
        all_models.update(ds.scores.keys())
    ordered = [m for m in MODEL_ORDER if m in all_models]
    tail = sorted(all_models - set(ordered))
    return ordered + tail


def _draw_tile_heatmap(
    matrix: np.ndarray,
    row_labels: list[str],
    col_labels: list[str],
    title: str,
    subtitle: str,
    output_path: Path,
    value_fmt: str = "{:.4f}",
    vmin: float = 0.0,
    vmax: float = 1.0,
) -> None:
    n_rows, n_cols = matrix.shape
    cmap = _custom_cmap()

    cell_w = 1.0
    cell_h = 0.9
    left_pad = 2.8
    top_pad = 1.6
    fig_w = left_pad + n_cols * cell_w + 0.6
    fig_h = top_pad + n_rows * cell_h + 0.8

    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    fig.patch.set_facecolor("#ffffff")
    ax.set_facecolor("#ffffff")
    ax.axis("off")

    # Title
    ax.text(
        0.0,
        top_pad + n_rows * cell_h + 0.45,
        title,
        fontsize=24,
        fontweight="bold",
        color="#18263E",
        ha="left",
        va="center",
    )
    if subtitle:
        ax.text(
            0.0,
            top_pad + n_rows * cell_h + 0.15,
            subtitle,
            fontsize=12,
            color="#5D6677",
            ha="left",
            va="center",
        )

    # Column labels
    for j, label in enumerate(col_labels):
        x = left_pad + j * cell_w + cell_w * 0.5
        ax.text(
            x,
            top_pad + n_rows * cell_h + 0.02,
            label,
            fontsize=12,
            fontweight="bold",
            color="#1C2940",
            ha="center",
            va="bottom",
        )

    for i, rlabel in enumerate(row_labels):
        y = top_pad + (n_rows - 1 - i) * cell_h + cell_h * 0.5
        ax.text(
            left_pad - 0.25,
            y,
            rlabel,
            fontsize=13,
            fontweight="bold",
            color="#202D45",
            ha="right",
            va="center",
        )

        for j in range(n_cols):
            x0 = left_pad + j * cell_w + 0.06
            y0 = top_pad + (n_rows - 1 - i) * cell_h + 0.08
            val = float(matrix[i, j])
            norm = 0.0 if vmax <= vmin else max(0.0, min(1.0, (val - vmin) / (vmax - vmin)))
            face = cmap(norm)
            patch = FancyBboxPatch(
                (x0, y0),
                cell_w - 0.12,
                cell_h - 0.16,
                boxstyle="round,pad=0.02,rounding_size=0.08",
                linewidth=0.0,
                facecolor=face,
                edgecolor="none",
            )
            ax.add_patch(patch)
            text_color = "#FFFFFF" if norm < 0.45 else "#0E2732"
            ax.text(
                x0 + (cell_w - 0.12) / 2,
                y0 + (cell_h - 0.16) / 2,
                value_fmt.format(val),
                fontsize=11.5,
                fontweight="bold",
                color=text_color,
                ha="center",
                va="center",
            )

    ax.set_xlim(-0.1, left_pad + n_cols * cell_w + 0.2)
    ax.set_ylim(0.0, top_pad + n_rows * cell_h + 0.8)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=260, bbox_inches="tight")
    plt.close(fig)


def _draw_grouped_bar(
    values_by_dataset: dict[str, list[float]],
    model_labels: list[str],
    metric_label: str,
    out_path: Path,
) -> None:
    datasets = list(values_by_dataset.keys())
    n_models = len(model_labels)
    x = np.arange(n_models)
    width = min(0.24, 0.76 / max(1, len(datasets)))

    palette = ["#3D7EA6", "#6CBF84", "#A27CBF", "#E4975D", "#5DA5A4", "#B6A84A"]

    fig, ax = plt.subplots(figsize=(12.8, 6.2))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    for i, ds in enumerate(datasets):
        vals = values_by_dataset[ds]
        offset = (i - (len(datasets) - 1) / 2.0) * width
        bars = ax.bar(x + offset, vals, width=width, label=ds.upper(), color=palette[i % len(palette)], alpha=0.92)
        for b in bars:
            h = b.get_height()
            ax.text(
                b.get_x() + b.get_width() / 2,
                h + 0.012,
                f"{h:.3f}",
                ha="center",
                va="bottom",
                fontsize=8,
                color="#334",
                rotation=90,
            )

    ax.set_title(f"{metric_label}: Cross-Dataset Model Comparison", fontsize=17, fontweight="bold", color="#1D2B45")
    ax.set_ylabel("Score", fontsize=11)
    ax.set_ylim(0.0, 1.05)
    ax.set_xticks(x)
    ax.set_xticklabels(model_labels, rotation=20, ha="right", fontsize=10)
    ax.grid(axis="y", linestyle="--", alpha=0.28)
    ax.legend(loc="upper right", ncol=len(datasets), frameon=False)
    fig.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=240, bbox_inches="tight")
    plt.close(fig)


def _build_matrix_for_dataset(ds: DatasetScore, model_ids: list[str], metrics: list[tuple[str, str]]) -> np.ndarray:
    mat = []
    for model_id in model_ids:
        row = []
        score_row = ds.scores.get(model_id, {})
        for key, _ in metrics:
            row.append(_as_float(score_row.get(key), 0.0))
        mat.append(row)
    return np.array(mat, dtype=float)


def _add_image_slide(prs: Presentation, title: str, image_path: Path, subtitle: str = "") -> None:
    layout = _pick_blank_layout(prs)
    slide = prs.slides.add_slide(layout)
    # title
    tx = slide.shapes.add_textbox(Inches(0.45), Inches(0.20), Inches(12.4), Inches(0.55))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.name = "Aptos"
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.name = "Aptos"

    slide.shapes.add_picture(str(image_path), Inches(0.35), Inches(0.90), width=Inches(12.65), height=Inches(6.35))


def _pick_blank_layout(prs: Presentation):
    # Be robust to custom templates with different layout counts/order.
    for layout in prs.slide_layouts:
        name = str(getattr(layout, "name", "") or "").lower()
        if "blank" in name:
            return layout
    return prs.slide_layouts[len(prs.slide_layouts) - 1]


def run(args: argparse.Namespace) -> Path:
    dataset_to_dir: dict[str, Path] = {}
    for pair in args.dataset_experiments:
        if "=" not in pair:
            raise ValueError(f"Invalid --dataset-experiments item: {pair}")
        ds, exp = pair.split("=", 1)
        ds = ds.strip()
        exp_dir = Path(exp).expanduser().resolve()
        if not exp_dir.exists():
            raise FileNotFoundError(f"Experiment dir not found for {ds}: {exp_dir}")
        dataset_to_dir[ds] = exp_dir

    ds_scores = [_read_primary_scores(exp_dir, ds) for ds, exp_dir in dataset_to_dir.items()]
    ds_scores.sort(key=lambda x: x.dataset_id)
    model_ids = _model_ids_present(ds_scores)
    model_labels = [MODEL_LABEL.get(m, m) for m in model_ids]

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    asset_dir = (args.asset_dir or (PROJECT_ROOT / "slides_output" / f"assets_eval_allsynsql_{ts}")).expanduser().resolve()
    asset_dir.mkdir(parents=True, exist_ok=True)

    # 1) Per-dataset main heatmaps (without overall)
    generated_images: list[tuple[str, Path, str]] = []
    for ds in ds_scores:
        mat = _build_matrix_for_dataset(ds, model_ids, METRICS_MAIN)
        cols = [label for _, label in METRICS_MAIN]
        out = asset_dir / f"heatmap_{ds.dataset_id}_main.png"
        _draw_tile_heatmap(
            mat,
            model_labels,
            cols,
            title=f"Model Capability Heatmap: {ds.dataset_id.upper()}",
            subtitle=f"Primary workload: {ds.primary_workload_run_id} (overall column removed)",
            output_path=out,
        )
        generated_images.append((f"{ds.dataset_id.upper()} Capability Heatmap", out, "Main metrics only"))

    # 2) Per-metric cross-dataset heatmaps + bars
    ds_ids = [d.dataset_id for d in ds_scores]
    for m_key, m_label in METRICS_ALL_SMALL:
        mat = []
        for model_id in model_ids:
            row = []
            for ds in ds_scores:
                row.append(_as_float(ds.scores.get(model_id, {}).get(m_key), 0.0))
            mat.append(row)
        mat_np = np.array(mat, dtype=float)
        out_heat = asset_dir / f"metric_heatmap_{m_key}.png"
        _draw_tile_heatmap(
            mat_np,
            model_labels,
            [x.upper() for x in ds_ids],
            title=f"{m_label}: Cross-Dataset Heatmap",
            subtitle="Rows=models, columns=datasets",
            output_path=out_heat,
        )
        generated_images.append((f"{m_label} (Heatmap)", out_heat, "Cross-dataset per-metric view"))

        # grouped bars
        values_by_dataset = {}
        for ds in ds_scores:
            values_by_dataset[ds.dataset_id] = [
                _as_float(ds.scores.get(model_id, {}).get(m_key), 0.0) for model_id in model_ids
            ]
        out_bar = asset_dir / f"metric_bars_{m_key}.png"
        _draw_grouped_bar(values_by_dataset, model_labels, m_label, out_bar)
        generated_images.append((f"{m_label} (Bars)", out_bar, "Cross-dataset per-metric bars"))

    # Export a compact numeric table for references
    summary_csv = asset_dir / "primary_scores_summary.csv"
    fieldnames = ["dataset_id", "primary_workload_run_id", "model_id"] + [k for k, _ in METRICS_ALL_SMALL] + ["overall_score"]
    with summary_csv.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for ds in ds_scores:
            for model_id in model_ids:
                rec = {
                    "dataset_id": ds.dataset_id,
                    "primary_workload_run_id": ds.primary_workload_run_id,
                    "model_id": model_id,
                }
                for mk, _ in METRICS_ALL_SMALL:
                    rec[mk] = _as_float(ds.scores.get(model_id, {}).get(mk), 0.0)
                rec["overall_score"] = _as_float(ds.scores.get(model_id, {}).get("overall_score"), 0.0)
                writer.writerow(rec)

    # 3) Append to PPT
    ppt_in = args.ppt_in.expanduser().resolve()
    ppt_out = args.ppt_out.expanduser().resolve()
    prs = Presentation(str(ppt_in))

    # Section opening slide
    slide = prs.slides.add_slide(_pick_blank_layout(prs))
    tx = slide.shapes.add_textbox(Inches(0.55), Inches(0.72), Inches(12.2), Inches(1.2))
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Appendix: Cross-Dataset Benchmark Evaluation (Syn + SQL)"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.name = "Aptos"
    p2 = tf.add_paragraph()
    p2.text = (
        "Datasets: " + ", ".join(x.upper() for x in ds_ids) +
        " | Main heatmaps exclude overall column."
    )
    p2.font.size = Pt(16)
    p2.font.name = "Aptos"
    p3 = tf.add_paragraph()
    p3.text = f"Generated at {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    p3.font.size = Pt(12)
    p3.font.name = "Aptos"

    for title, img, subtitle in generated_images:
        _add_image_slide(prs, title, img, subtitle)

    prs.save(str(ppt_out))

    manifest = {
        "timestamp": datetime.now().isoformat(),
        "datasets": ds_ids,
        "experiments": {d.dataset_id: str(d.experiment_dir) for d in ds_scores},
        "primary_workloads": {d.dataset_id: d.primary_workload_run_id for d in ds_scores},
        "asset_dir": str(asset_dir),
        "summary_csv": str(summary_csv),
        "ppt_in": str(ppt_in),
        "ppt_out": str(ppt_out),
        "slides_appended": len(generated_images) + 1,
    }
    (asset_dir / "append_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return ppt_out


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate cross-dataset evaluation visuals and append to PPT.")
    parser.add_argument(
        "--dataset-experiments",
        nargs="+",
        required=True,
        help="Pairs in form dataset_id=/abs/path/to/experiment_dir",
    )
    parser.add_argument(
        "--ppt-in",
        type=Path,
        default=PROJECT_ROOT / "slides_output" / "v4.1.pptx",
    )
    parser.add_argument(
        "--ppt-out",
        type=Path,
        default=PROJECT_ROOT / "slides_output" / "v4.2_all_syn_sql_eval_append.pptx",
    )
    parser.add_argument(
        "--asset-dir",
        type=Path,
        default=None,
        help="Optional output dir for generated PNG assets.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
